"""OOT 트렌드 리포트 엑셀 생성 — 첨부 이미지 양식 재현(참조 구현).

검사항목별로 시트 1개: 기준정보 블록 + 통계 블록 + LOT 데이터표 + 관리도 차트.
계산셀(평균·표준편차·Max·min·평균±3σ·Cpk·기준일탈여부)은 데이터표를 참조하는 수식.
데이터표(제조번호·Lot평균·최대값·최소값)만 채우면 나머지는 자동 계산된다.
"""
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "맑은 고딕"
GREEN = "C4D79B"      # 라벨 셀
GREEN_LT = "EBF1DE"   # 값 셀(옅은 초록)
NUM6 = "0.000000"
NUM2 = "0.00"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _label(cell):
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.font = Font(name=FONT, size=9, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER


def _value(cell, num_fmt=None, light=True):
    if light:
        cell.fill = PatternFill("solid", fgColor=GREEN_LT)
    cell.font = Font(name=FONT, size=9)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER
    if num_fmt:
        cell.number_format = num_fmt


def build_report(wb, cfg):
    ws = wb.create_sheet(cfg["sheet"])
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 12, "B": 14, "C": 12, "D": 12, "E": 14, "F": 14, "G": 12, "H": 12}.items():
        ws.column_dimensions[col].width = w

    n = len(cfg["values"])
    first, last = 16, 15 + n  # 데이터 행 범위
    data_avg = f"$B${first}:$B${last}"

    # ---- 규격 유형별 Cp/K/Cpk·기준일탈 (B7=USL 상한/이하, B8=LSL 하한/이상) ----
    # parse_criterion 결과를 cfg['규격하한'](LSL,이상)·cfg['규격상한'](USL,이하)로 받는다.
    # 구버전 cfg(자가기준상한/하한만) 호환: 규격키가 없으면 자가기준하한=LSL, 값만 있으면 이상(하한)규격으로 간주.
    lo = cfg.get("규격하한", cfg.get("자가기준하한"))   # LSL(이상)
    hi = cfg.get("규격상한")                            # USL(이하)
    if "규격하한" not in cfg and "규격상한" not in cfg:  # 구버전 호환(이상 규격 가정)
        lo = cfg.get("자가기준상한"); hi = None
    has_lo, has_hi = lo is not None, hi is not None
    if has_lo and has_hi:            # 범위 규격 → 표준 Cp/K/Cpk (점검표 템플릿과 동일)
        f_cp = "=(B7-B8)/(6*H4)"                       # Cp=(USL-LSL)/6σ
        f_k = "=ABS((B7+B8)/2-H3)/((B7-B8)/2)"         # K=|중앙값-μ|/반폭 (절대값)
        f_cpk = "=IF(H8=0,H7,IF(H8>=1,0,(1-H8)*H7))"   # Cpk=(1-K)·Cp (=min(CPU,CPL))
        f_dev = '=IF(AND(H5<=B7,H6>=B8),"기준일탈없음","기준일탈있음")'
    elif has_hi:                     # 상한(이하)만 → Cpk=CPU=(USL-μ)/3σ
        f_cp = f_k = "N/A"
        f_cpk = "=(B7-H3)/(3*H4)"
        f_dev = '=IF(H5<=B7,"기준일탈없음","기준일탈있음")'
    elif has_lo:                     # 하한(이상)만 → Cpk=CPL=(μ-LSL)/3σ
        f_cp = f_k = "N/A"
        f_cpk = "=(H3-B8)/(3*H4)"
        f_dev = '=IF(H6>=B8,"기준일탈없음","기준일탈있음")'
    else:                            # 규격 없음(정성/미인식) → 전부 N/A
        f_cp = f_k = f_cpk = f_dev = "N/A"

    # ---- 제목 ----
    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = cfg["title"]
    t.font = Font(name=FONT, size=13, bold=True)
    t.alignment = Alignment(horizontal="center", vertical="center")
    t.fill = PatternFill("solid", fgColor=GREEN)
    ws.row_dimensions[1].height = 24

    # ---- 기준정보 블록 (A3:B13) ----
    info = [
        ("제품명", cfg["제품명"], None),
        ("제조공정", cfg["제조공정"], None),
        ("검사항목", cfg["검사항목"], None),
        ("허가기준", cfg["허가기준"], None),
        ("자가기준상한(USL·이하)", hi if has_hi else "", None),   # B7 = 상한(USL)
        ("자가기준하한(LSL·이상)", lo if has_lo else "", None),   # B8 = 하한(LSL)
        ("평균 +3σ", "=H3+3*H4", NUM6),
        ("평균 -3σ", "=H3-3*H4", NUM6),
        ("단위", cfg["단위"], None),
        ("자가기준 (1차)", cfg["자가기준_1차"], None),
        ("자가기준 (2차)", cfg["자가기준_2차"], None),
    ]
    for i, (lab, val, fmt) in enumerate(info):
        r = 3 + i
        ws[f"A{r}"] = lab
        _label(ws[f"A{r}"])
        ws[f"B{r}"] = val if val is not None else ""
        _value(ws[f"B{r}"], fmt)

    # ---- 통계 블록 (G3:H10) ----
    stats = [
        ("평균", f"=AVERAGE({data_avg})", NUM6),
        ("표준편차", f"=STDEV({data_avg})", NUM2),
        ("Max", f"=MAX({data_avg})", NUM6),
        ("min", f"=MIN({data_avg})", NUM6),
        ("Cp", f_cp, NUM2 if f_cp.startswith("=") else None),
        ("K", f_k, NUM2 if f_k.startswith("=") else None),
        ("Cpk", f_cpk, NUM2 if f_cpk.startswith("=") else None),
        ("기준일탈여부", f_dev, None),
    ]
    for i, (lab, val, fmt) in enumerate(stats):
        r = 3 + i
        ws[f"G{r}"] = lab
        _label(ws[f"G{r}"])
        ws[f"H{r}"] = val
        _value(ws[f"H{r}"], fmt)

    # ---- 데이터표 헤더 (row 15) ----
    headers = ["제조번호", "Lot 평균", "최대값", "최소값",
               "기준하한(이상)", "기준상한(이하)", "평균 -3σ", "평균 +3σ"]
    for j, h in enumerate(headers):
        c = ws.cell(row=15, column=1 + j, value=h)
        _label(c)

    # ---- 데이터 행 ----
    for i in range(n):
        r = first + i
        ws.cell(row=r, column=1, value=cfg["lots"][i])                    # 제조번호
        _value(ws.cell(row=r, column=1), light=False)
        ws.cell(row=r, column=2, value=cfg["values"][i]).number_format = NUM6  # Lot 평균
        _value(ws.cell(row=r, column=2), NUM6, light=False)
        for col in (3, 4):                                                # 최대값·최소값 (입력 시 채움)
            _value(ws.cell(row=r, column=col), NUM6, light=False)
        ws.cell(row=r, column=5, value=("=$B$8" if has_lo else "N/A"))    # 기준하한(이상)=LSL
        _value(ws.cell(row=r, column=5), NUM6 if has_lo else None, light=False)
        ws.cell(row=r, column=6, value=("=$B$7" if has_hi else "N/A"))    # 기준상한(이하)=USL
        _value(ws.cell(row=r, column=6), NUM6 if has_hi else None, light=False)
        ws.cell(row=r, column=7, value="=$B$10")                          # 평균 -3σ
        _value(ws.cell(row=r, column=7), NUM6, light=False)
        ws.cell(row=r, column=8, value="=$B$9")                           # 평균 +3σ -> B9
        _value(ws.cell(row=r, column=8), NUM6, light=False)

    # ---- 관리도 차트 ----
    chart = LineChart()
    chart.title = cfg["title"]
    chart.style = 2
    chart.y_axis.title = cfg["단위"]
    chart.x_axis.title = "제조번호"
    chart.height = 9
    chart.width = 26

    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    series_def = [
        (4, "4472C4", 1.5, None),      # 최소값 (blue)
        (2, "FF0000", 2.5, None),      # Lot 평균 (red, thick)
        (3, "00B0A0", 1.5, None),      # 최대값 (teal)
        (5, "C00000", 1.25, "dash"),   # 기준하한(이상, LSL)
        (6, "C00000", 1.25, "dash"),   # 기준상한(이하, USL)
        (7, "FFC000", 1.25, "dash"),   # 평균 -3σ
        (8, "FFC000", 1.25, "dash"),   # 평균 +3σ
    ]
    for col, color, width, dash in series_def:
        ref = Reference(ws, min_col=col, min_row=15, max_row=last)  # header 포함
        chart.add_data(ref, titles_from_data=True)
        s = chart.series[-1]
        lp = LineProperties(solidFill=color, w=int(width * 12700))
        if dash:
            lp.prstDash = dash
        s.graphicalProperties = GraphicalProperties()
        s.graphicalProperties.line = lp
        s.smooth = False

    chart.x_axis.delete = False
    chart.y_axis.delete = False
    ws.add_chart(chart, f"A{last + 3}")
    return ws


PAEONI = {
    "sheet": "패오니플로린",
    "title": "광동원탕 완제품시험 작약 중 패오니플로린",
    "제품명": "광동원탕", "제조공정": "완제품시험", "검사항목": "작약 중 패오니플로린",
    "허가기준": "13.1 mg/100mL이상", "자가기준상한": 13.1, "자가기준하한": None,
    "단위": "mg/100ml", "자가기준_1차": "13.1 mg/100mL 이상", "자가기준_2차": None,
    "lots": [f"250{i:02d}" for i in range(1, 45)],
    "values": [35.9, 41.1, 41.2, 41.1, 40.4, 41.5, 42.0, 38.0, 38.0, 37.9, 38.9, 40.4,
               37.7, 38.3, 36.5, 32.3, 32.3, 32.0, 31.0, 31.1, 32.7, 31.5, 33.2, 30.9,
               28.9, 26.4, 26.5, 26.4, 26.2, 26.0, 26.7, 26.6, 25.3, 25.1, 25.7, 25.7,
               25.7, 26.2, 26.2, 26.0, 27.6, 28.5, 30.6, 30.5],
}
GLYCYRRHIZIC = {
    "sheet": "글리시리진산",
    "title": "광동원탕 완제품시험 감초 중 글리시리진산",
    "제품명": "광동원탕", "제조공정": "완제품시험", "검사항목": "감초 중 글리시리진산",
    "허가기준": "4.7 mg/100mL이상", "자가기준상한": 4.7, "자가기준하한": None,
    "단위": "mg/100ml", "자가기준_1차": "4.7 mg/100mL 이상", "자가기준_2차": None,
    "lots": [f"250{i:02d}" for i in range(1, 45)],
    "values": [10.9, 11.3, 11.4, 11.2, 10.6, 10.8, 10.9, 11.2, 11.3, 11.3, 11.0, 11.1,
               9.2, 9.3, 9.3, 9.9, 9.9, 9.4, 9.0, 9.0, 9.0, 9.3, 9.1, 8.6, 8.9, 8.8,
               9.0, 9.0, 9.8, 9.5, 9.4, 9.4, 9.8, 9.5, 9.5, 9.2, 8.4, 9.0, 9.2, 9.5,
               9.3, 9.7, 9.9, 10.0],
}

if __name__ == "__main__":
    wb = Workbook()
    wb.remove(wb.active)
    build_report(wb, PAEONI)
    build_report(wb, GLYCYRRHIZIC)
    wb.save("sample_OOT_report.xlsx")
    print("saved")
