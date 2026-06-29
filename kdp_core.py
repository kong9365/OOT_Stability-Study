# -*- coding: utf-8 -*-
"""
kdp_core.py — 광동 품질·시험 통합 대시보드 백엔드 로직 (Streamlit 비의존).

FastAPI(dashboard_api.py)와 필요 시 다른 진입점이 공유하는 순수 Python 모듈.
Tableau OOT_추출용 뷰에서 데이터를 받아 OOT 판정 / 안정성 회귀 / ANCOVA를 계산한다.

핵심 함수
  - oot_products(test_type)            : (품목코드, 품목) 목록
  - oot_lot_summary(code, test_type)   : LOT별 정상/주의/관리이탈/정성 카운트 + OOT 항목
  - stability_analysis(code, ...)      : 시점별 회귀·유효기간·시점간 OOT·ANCOVA
"""
from __future__ import annotations

import csv
import io
import os
import re
import threading
import time
from typing import Optional

import httpx
import numpy as np
import pandas as pd
import urllib3
from dotenv import load_dotenv

from stability import analyze_dataframe, ci_bound  # 검증된 회귀 엔진 재사용

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"), override=False)

TABLEAU_SERVER = os.getenv("TABLEAU_SERVER", "http://tableau.ekdp.com").rstrip("/")
API_VER = os.getenv("TABLEAU_API_VERSION", "3.21")
PAT_NAME = os.getenv("TABLEAU_PAT_NAME", "MISO")
PAT_SECRET = os.getenv("TABLEAU_PAT_SECRET", "")
BASE = f"{TABLEAU_SERVER}/api/{API_VER}"
OOT_WORKBOOK = "OOT 대시보드"
OOT_VIEW_NAME = "OOT_추출용"

# 시험종류 — Tableau vf_시험종류 서버측 필터값(완제품 기본). 안정성 종류는 '안정성' 포함.
OOT_TEST_TYPES = [
    "완제품", "원료시험", "반제품", "직접자재", "시험기기 및 기구", "기타시험",
    "장기 안정성시험(Long Term)", "가속 안정성시험(Acclerated)",
    "시판후 안정성시험(Ongoing Stability)", "4b장기 안정성시험(LT4b)",
]
STAB_TEST_TYPES = [t for t in OOT_TEST_TYPES if "안정성" in t]

# 안정성 차트 배치 색상 팔레트(디자인 일치)
BATCH_COLORS = ["#E5310F", "#1F7A52", "#B8893B", "#2563eb", "#9333ea", "#0891b2"]

# ── Tableau 클라이언트 (토큰·뷰ID 캐시) ─────────────────────────────────────
_LOCK = threading.Lock()
_client = httpx.Client(timeout=120.0, verify=False)
_auth = {"token": None, "site": None, "view_id": None, "ts": 0.0}


def _signin() -> tuple[str, str]:
    if not PAT_SECRET:
        raise RuntimeError("TABLEAU_PAT_SECRET 미설정(.env)")
    r = _client.post(f"{BASE}/auth/signin", json={"credentials": {
        "personalAccessTokenName": PAT_NAME, "personalAccessTokenSecret": PAT_SECRET,
        "site": {"contentUrl": ""}}},
        headers={"Accept": "application/json", "Content-Type": "application/json"})
    if r.status_code != 200:
        raise RuntimeError(f"Tableau 인증 실패 {r.status_code}: {r.text[:160]}")
    cr = r.json()["credentials"]
    return cr["token"], cr["site"]["id"]


def _resolve_view_id(tok: str, site: str) -> Optional[str]:
    H = {"X-Tableau-Auth": tok, "Accept": "application/json"}
    wbm = {w["id"]: w["name"] for w in _client.get(
        f"{BASE}/sites/{site}/workbooks", headers=H, params={"pageSize": "1000"}
    ).json().get("workbooks", {}).get("workbook", [])}
    for v in _client.get(f"{BASE}/sites/{site}/views", headers=H,
                         params={"pageSize": "1000"}).json().get("views", {}).get("view", []):
        if v["name"] == OOT_VIEW_NAME and wbm.get(v.get("workbook", {}).get("id", "")) == OOT_WORKBOOK:
            return v["id"]
    return None


def _ensure() -> tuple[str, str, str]:
    """토큰·사이트·뷰ID 확보(15분 캐시, 만료/401시 재인증). 스레드 안전."""
    with _LOCK:
        if _auth["token"] and (time.monotonic() - _auth["ts"] < 900):
            return _auth["token"], _auth["site"], _auth["view_id"]
        tok, site = _signin()
        vid = _resolve_view_id(tok, site)
        if not vid:
            raise RuntimeError("OOT_추출용 뷰를 찾지 못함")
        _auth.update(token=tok, site=site, view_id=vid, ts=time.monotonic())
        return tok, site, vid


def _view_csv(params: dict) -> str:
    tok, site, vid = _ensure()
    url = f"{BASE}/sites/{site}/views/{vid}/data"
    r = _client.get(url, headers={"X-Tableau-Auth": tok}, params=params)
    if r.status_code == 401:                       # 토큰 만료 → 강제 재인증 1회
        with _LOCK:
            _auth["token"] = None
        tok, site, vid = _ensure()
        url = f"{BASE}/sites/{site}/views/{vid}/data"
        r = _client.get(url, headers={"X-Tableau-Auth": tok}, params=params)
    r.raise_for_status()
    return r.text


# ── 유틸 ─────────────────────────────────────────────────────────────────────
def _num(x) -> float:
    try:
        return float(str(x).replace(",", "").strip())
    except (ValueError, TypeError):
        return float("nan")


def _oot_status(s) -> str:
    s = str(s or "")
    if "주의" in s:                 # '주의 (±2σ~±3σ)'에 3σ가 들어있어 주의를 먼저 판정
        return "주의"
    if "관리이탈" in s or "초과" in s:
        return "관리이탈"
    if "정상" in s:
        return "정상"
    return "데이터부족"


_QUAL_KW = ("불검출", "성상", "현탁액", "반점", "Rf", "S/N", "일치", "검액",
            "음성", "양성", "적합하다", "투명", "색상")
_UNIT = r"(?:mg|㎎|μg|㎍|ug|g|%|ppm)"


def parse_criterion(text):
    """시험기준 텍스트 → (규격하한, 규격상한). stability_page.parse_criterion 과 동일 규칙."""
    if text is None or (isinstance(text, float) and np.isnan(text)):
        return (None, None)
    s = str(text).strip()
    if not s or any(k in s for k in _QUAL_KW):
        return (None, None)
    m = re.search(r"(\d+\.?\d*)\s*[∼~〜–—\-]\s*(\d+\.?\d*)", s)
    if m:
        return (float(m.group(1)), float(m.group(2)))
    m = re.search(r"(\d+\.?\d*)\s*" + _UNIT, s)
    if not m:
        return (None, None)
    val = float(m.group(1))
    if "이하" in s or "미만" in s:
        return (None, val)
    if "이상" in s or "초과" in s:
        return (val, None)
    return (None, None)


def _year_of(lot: str) -> str:
    return ("20" + lot[:2]) if lot[:2].isdigit() else "기타"


def _parse_date(s):
    """'2021. 8. 13.' / '2024-08-12' → date. 실패 시 None."""
    from datetime import date
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", str(s or ""))
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _approved_months(mfg, exp):
    """허가 유효기간(개월) = 유효기한 − 제조일자. 가장 가까운 정수 개월로 반올림."""
    a, b = _parse_date(mfg), _parse_date(exp)
    if not a or not b:
        return None
    days = (b - a).days
    return int(round(days / 30.4375)) if days > 0 else None


# ── OOT ──────────────────────────────────────────────────────────────────────
def oot_products(test_type: str = "완제품") -> list[dict]:
    """(품목코드, 품목) 고유 목록. 2열만 추출(경량)."""
    text = _view_csv({"vf_시험종류": test_type})
    seen: dict[str, str] = {}
    for row in csv.DictReader(io.StringIO(text)):
        code = (row.get("품목코드") or "").strip()
        if code and code not in seen:
            seen[code] = (row.get("품목") or "").strip()
    return [{"code": k, "name": v} for k, v in sorted(seen.items())]


def _product_df(code: str, test_type: str) -> pd.DataFrame:
    text = _view_csv({"vf_품목코드": str(code), "vf_시험종류": test_type})
    return pd.DataFrame(list(csv.DictReader(io.StringIO(text))))


def oot_lot_summary(code: str, test_type: str = "완제품") -> dict:
    """선택 품목의 LOT별 판정 요약. 한 번의 Tableau 호출로 전 LOT 계산.

    반환: {code, name, lots:[{lot, year, normal, warn, crit, qual, flagged,
            items:[{name, val, mean, sd, z, status}]}], years:[...]}
    """
    pdf = _product_df(code, test_type)
    if pdf.empty or "제조번호" not in pdf.columns:
        return {"code": code, "name": "", "lots": [], "years": ["전체"]}

    name = str(pdf.get("품목", pd.Series([""])).iloc[0]) if "품목" in pdf else ""
    val_col = next((c for c in ["LOT결과_0제외", "시험결과_유효", "평균 시험결과"] if c in pdf.columns), None)
    pdf["_상태"] = pdf["OOT_구간분류_히트맵"].map(_oot_status)
    pdf["_sd"] = pdf["확인_표준편차"].map(_num)
    pdf["_mu"] = pdf["확인_평균"].map(_num)
    pdf["_v"] = pdf[val_col].map(_num) if val_col else float("nan")

    lots = []
    for lot, g in pdf.groupby(pdf["제조번호"].astype(str)):
        lot = str(lot).strip()
        if not lot:
            continue
        judged = g[g["_sd"] > 0]                       # 정성/데이터부족(σ 0·NULL) 제외
        crit = int((judged["_상태"] == "관리이탈").sum())
        warn = int((judged["_상태"] == "주의").sum())
        normal = int((judged["_상태"] == "정상").sum())
        qual = int(len(g) - len(judged))
        def _item(r):
            z = (r["_v"] - r["_mu"]) / r["_sd"] if r["_sd"] else float("nan")
            return {
                "name": str(r.get("시험항목", "")), "status": r["_상태"],
                "val": None if np.isnan(r["_v"]) else round(float(r["_v"]), 4),
                "mean": None if np.isnan(r["_mu"]) else round(float(r["_mu"]), 4),
                "sd": None if np.isnan(r["_sd"]) else round(float(r["_sd"]), 4),
                "z": None if np.isnan(z) else round(float(z), 3),
            }
        items = [_item(r) for _, r in judged[judged["_상태"].isin(["주의", "관리이탈"])].iterrows()]
        items.sort(key=lambda it: ({"관리이탈": 2, "주의": 1}.get(it["status"], 0),
                                   abs(it["z"]) if it["z"] is not None else 0), reverse=True)
        normal_items = [_item(r) for _, r in judged[judged["_상태"] == "정상"].iterrows()]
        normal_items.sort(key=lambda it: abs(it["z"]) if it["z"] is not None else 0, reverse=True)
        lots.append({"lot": lot, "year": _year_of(lot), "normal": normal, "warn": warn,
                     "crit": crit, "qual": qual, "flagged": bool(items),
                     "items": items, "normalItems": normal_items})

    lots.sort(key=lambda l: l["lot"], reverse=True)
    years = sorted({l["year"] for l in lots if l["year"] != "기타"}, reverse=True)
    if any(l["year"] == "기타" for l in lots):
        years.append("기타")
    return {"code": code, "name": name, "lots": lots, "years": ["전체"] + years}


# ── 안정성 ────────────────────────────────────────────────────────────────────
def stability_df(code: str, test_type: str) -> pd.DataFrame:
    """OOT_추출용 → 안정성 long format(제조번호·시험항목·대분류·시점(개월)·결과값·규격)."""
    from lims_parser import parse_month
    raw = _product_df(code, test_type)
    if raw.empty:
        return pd.DataFrame()
    rec = []
    for _, row in raw.iterrows():
        month = parse_month(row.get("의뢰 특이사항", ""), test_type)
        val = _num(row.get("LOT결과_0제외", ""))
        if month is None or np.isnan(val):
            continue
        lo, hi = parse_criterion(row.get("시험기준", ""))
        rec.append({"제조번호": str(row.get("제조번호", "")).strip(),
                    "시험항목": str(row.get("시험항목", "")).strip(),
                    "대분류": str(row.get("대분류", "")).strip(),
                    "시점(개월)": month, "결과값": val,
                    "규격하한": lo if lo is not None else float("nan"),
                    "규격상한": hi if hi is not None else float("nan"),
                    "제조일자": str(row.get("제조일자", "")).strip(),
                    "유효기한": str(row.get("유효기한", "")).strip()})

    # 0개월 기준값 = 완제품(출하시험) 결과 포함 (ICH·원 LIMS 앱과 동일: 출하시점 = 0개월).
    # 안정성 배치에 매칭되는 출하결과를 0개월 점으로 추가 → 시점 1개 늘어 회귀 가능해짐.
    if "안정성" in test_type and rec:
        try:
            rel = _product_df(code, "완제품")
        except Exception:
            rel = pd.DataFrame()
        if not rel.empty and "제조번호" in rel.columns:
            stab_keys = {(r["제조번호"], r["시험항목"]) for r in rec}
            have0 = {(r["제조번호"], r["시험항목"]) for r in rec if r["시점(개월)"] == 0.0}
            for _, row in rel.iterrows():
                b = str(row.get("제조번호", "")).strip()
                item = str(row.get("시험항목", "")).strip()
                val = _num(row.get("LOT결과_0제외", ""))
                key = (b, item)
                if key not in stab_keys or key in have0 or np.isnan(val):
                    continue
                lo, hi = parse_criterion(row.get("시험기준", ""))
                rec.append({"제조번호": b, "시험항목": item,
                            "대분류": str(row.get("대분류", "")).strip(),
                            "시점(개월)": 0.0, "결과값": val,
                            "규격하한": lo if lo is not None else float("nan"),
                            "규격상한": hi if hi is not None else float("nan"),
                            "제조일자": str(row.get("제조일자", "")).strip(),
                            "유효기한": str(row.get("유효기한", "")).strip()})
                have0.add(key)
    return pd.DataFrame(rec)


def _ancova(d: pd.DataFrame) -> Optional[dict]:
    """함량 ~ 시점 + C(제조번호) + 시점:C(제조번호) 순차(Type I) ANOVA.

    교호작용 p < 0.25 → 배치별 개별 회귀, p ≥ 0.25 → 합산 가능(ICH Q1E·α=0.25).
    배치 2개 이상·각 배치 시점 2개 이상일 때만 산출. 불가 시 None.
    """
    bc = d.groupby("제조번호")["시점(개월)"].nunique()
    good = bc[bc >= 2].index
    d = d[d["제조번호"].isin(good)]
    nb = d["제조번호"].nunique()
    # 완전교호작용 모형 파라미터 = 2·배치수. 잔차자유도 ≥1 필요 → obs > 2·배치수.
    if nb < 2 or len(d) <= nb * 2:
        return None
    try:
        import statsmodels.formula.api as smf
        import statsmodels.api as sm
        dd = d.rename(columns={"결과값": "y", "시점(개월)": "t", "제조번호": "b"}).copy()
        dd["b"] = dd["b"].astype(str)
        model = smf.ols("y ~ t + C(b) + t:C(b)", data=dd).fit()
        tbl = sm.stats.anova_lm(model, typ=1)
    except Exception:
        return None

    label = {"t": "시간 (Month)", "C(b)": "제조번호 (Batch)",
             "t:C(b)": "시간×제조번호", "Residual": "오차 (Error)"}
    rows, inter_p = [], None
    for src in ["t", "C(b)", "t:C(b)", "Residual"]:
        if src not in tbl.index:
            continue
        r = tbl.loc[src]
        f = r.get("F"); p = r.get("PR(>F)")
        is_inter = src == "t:C(b)"
        if is_inter and pd.notna(p):
            inter_p = float(p)
        rows.append({
            "source": label[src], "df": int(r["df"]),
            "ss": round(float(r["sum_sq"]), 2),
            "f": None if (pd.isna(f)) else round(float(f), 2),
            "p": None if (pd.isna(p)) else round(float(p), 3),
            "interaction": is_inter, "error": src == "Residual",
        })
    if inter_p is None:
        return None
    poolable = inter_p >= 0.25
    return {
        "alpha": 0.25, "rows": rows, "interactionP": round(inter_p, 3),
        "poolable": poolable,
        "verdict": "합산(pooled) 가능" if poolable else "배치별 개별 회귀",
        "note": ("교호작용(시간×제조번호) p ≥ 0.25 → 기울기 공통. 제조번호 주효과로 절편 공통 여부를 "
                 "판정해 합산 회귀를 적용합니다."
                 if poolable else
                 f"교호작용(시간×제조번호) p = {inter_p:.3f} < 0.25 → 배치 간 분해속도가 달라 합산 불가. "
                 "배치별 개별 회귀선으로 각각 유효기간을 산출하고 최단값(worst-case)을 대표값으로 채택합니다."),
    }


def _pick_test_item(df: pd.DataFrame) -> Optional[str]:
    """기본 시험항목 선택 — 함량(대분류) 항목 우선, 회귀가능 데이터 많은 순."""
    best, best_key = None, (-1, -1)
    for item, g in df.groupby("시험항목"):
        pairs = g.groupby("제조번호")["시점(개월)"].nunique()
        n = int(pairs[pairs >= 2].sum())
        is_assay = 1 if g["대분류"].astype(str).str.contains("함량").any() else 0
        key = (is_assay, n)              # 함량 우선, 그다음 데이터 많은 항목
        if key > best_key:
            best, best_key = item, key
    return best if best_key[1] >= 2 else (df["시험항목"].iloc[0] if len(df) else None)


def _steps_json(steps: dict) -> list:
    """RegressionResult.steps(중간계산) → JSON 안전한 [{k, v}] 순서 보존 리스트."""
    out = []
    for k, v in steps.items():
        if isinstance(v, float):
            out.append({"k": k, "v": None if (np.isnan(v)) else round(float(v), 6)})
        else:
            out.append({"k": k, "v": v})
    return out


def stability_analysis(code: str, test_type: str, spec_low: float = 90.0,
                       spec_high: float = 150.0, method: str = "pooled",
                       test_item: Optional[str] = None, batch: Optional[str] = None) -> dict:
    """단일 시험항목에 대한 회귀·유효기간·시점간 OOT·ANCOVA 결과(JSON 친화 dict).

    batch 지정 시 해당 제조번호(LOT) 한 배치만 분석한다(OOT 교차연동용).
    """
    df = stability_df(code, test_type)
    if df.empty:
        return {"ok": False, "reason": "정량 시점 데이터 없음"}
    if batch:
        df = df[df["제조번호"].astype(str) == str(batch)]
        if df.empty:
            return {"ok": False, "reason": f"배치 {batch}의 정량 시점 데이터가 없습니다."}

    items = sorted(df["시험항목"].unique())
    item = test_item if test_item in items else _pick_test_item(df)
    if not item:
        return {"ok": False, "reason": "분석 가능한 시험항목 없음"}
    d = df[df["시험항목"] == item].copy()

    # 규격: 시험기준 파싱값 우선, 양쪽 모두 없으면 수동 기본값
    both_na = d["규격하한"].isna() & d["규격상한"].isna()
    d.loc[both_na, "규격하한"] = spec_low
    d.loc[both_na, "규격상한"] = spec_high
    sl = float(d["규격하한"].dropna().iloc[0]) if d["규격하한"].notna().any() else None
    sh = float(d["규격상한"].dropna().iloc[0]) if d["규격상한"].notna().any() else None

    # 허가 유효기간(개월) = 유효기한 − 제조일자 (배치별, 행정상 승인값)
    approved = {}
    for b, gb in d.groupby("제조번호"):
        mfg = next((x for x in gb.get("제조일자", []) if str(x).strip()), "")
        exp = next((x for x in gb.get("유효기한", []) if str(x).strip()), "")
        approved[str(b)] = _approved_months(mfg, exp)

    col_map = {"batch": "제조번호", "test": "시험항목", "major": "대분류",
               "time": "시점(개월)", "assay": "결과값",
               "spec_low": "규격하한", "spec_high": "규격상한"}
    results = analyze_dataframe(d, col_map, spec_low, spec_high, method=method)
    if not results:
        return {"ok": False, "reason": "회귀 가능한 데이터 없음(시점 2개 이상 필요)"}

    batches = []
    for i, r in enumerate(sorted(results, key=lambda r: r.batch)):
        batches.append({
            "batch": r.batch, "color": BATCH_COLORS[i % len(BATCH_COLORS)],
            "approved": approved.get(str(r.batch)),
            "pts": [[float(x), float(y)] for x, y in zip(r.x, r.y)],
            "slope": None if np.isnan(r.slope) else round(float(r.slope), 5),
            "intercept": None if np.isnan(r.intercept) else round(float(r.intercept), 3),
            "r2": None if np.isnan(r.r_squared) else round(float(r.r_squared), 4),
            "n": int(r.n), "se": None if np.isnan(r.se_resid) else float(r.se_resid),
            "tcrit": None if np.isnan(r.t_crit) else float(r.t_crit),
            "mx": float(np.mean(r.x)) if r.n else 0.0,
            "sxx": float(np.sum((r.x - np.mean(r.x)) ** 2)) if r.n else 0.0,
            "shelf": None if r.shelf_life is None else round(float(r.shelf_life), 1),
            "trend": r.trend, "lastT": float(r.last_timepoint),
            "limit": r.limiting_side, "steps": _steps_json(r.steps),
        })

    valid = [b for b in batches if b["shelf"] is not None]
    worst = min(valid, key=lambda b: b["shelf"]) if valid else None

    # 시점간 OOT — worst(없으면 첫) 배치 기준
    tb_name = worst["batch"] if worst else batches[0]["batch"]
    tb = next((r for r in results if r.batch == tb_name), results[0])
    sigma = float(tb.se_resid) if (tb.se_resid and not np.isnan(tb.se_resid)) else float("nan")
    tp_rows, n_oot, n_warn = [], 0, 0
    if not np.isnan(sigma) and sigma > 0 and not np.isnan(tb.slope):
        for x, y in zip(tb.x, tb.y):
            pred = tb.intercept + tb.slope * x
            dev = (y - pred) / sigma
            ad = abs(dev)
            judge = "⚠ OOT (3σ 밖)" if ad > 3 else ("주의 (2~3σ)" if ad > 2 else "정상 (2σ 이내)")
            if ad > 3:
                n_oot += 1
            elif ad > 2:
                n_warn += 1
            tp_rows.append({"time": round(float(x), 0), "meas": round(float(y), 2),
                            "pred": round(float(pred), 2), "dev": round(float(dev), 2),
                            "judge": judge})

    ancova = _ancova(d)
    pair_count = int(sum(len(b["pts"]) for b in batches))
    # 허가 유효기간 대표값(최빈값) + 추정 저장수명이 허가보다 짧은지 플래그
    av = [v for v in approved.values() if v]
    approved_months = max(set(av), key=av.count) if av else None
    approved_short = bool(worst and approved_months and worst["shelf"] < approved_months)
    return {
        "ok": True, "code": code, "testType": test_type, "testItem": item,
        "testItems": items, "batch": batch, "method": method, "specLow": sl, "specHigh": sh,
        "batchCount": len(batches), "pairCount": pair_count,
        "approvedMonths": approved_months, "approvedShort": approved_short,
        "batches": batches, "worst": worst,
        "sigma": None if np.isnan(sigma) else round(sigma, 3),
        "sigmaBase": "통합 잔차오차" if method == "pooled" else "배치 잔차오차",
        "tpBatch": tb_name, "tpRows": tp_rows, "nOot": n_oot, "nWarn": n_warn,
        "ancova": ancova,
    }


def stability_tables(code: str, test_type: str, spec_low: float = 90.0,
                     spec_high: float = 150.0, batch: Optional[str] = None) -> dict:
    """전 시험항목 × 배치에 대한 두 방식 비교표·상세 요약·원자료 (기존 LIMS 앱 정보 복원).

    batch 지정 시 해당 배치만.
    """
    df = stability_df(code, test_type)
    if df.empty:
        return {"ok": False, "reason": "데이터 없음"}
    df = df.copy()
    if batch:
        df = df[df["제조번호"].astype(str) == str(batch)]
        if df.empty:
            return {"ok": False, "reason": f"배치 {batch} 데이터 없음"}
    both_na = df["규격하한"].isna() & df["규격상한"].isna()
    df.loc[both_na, "규격하한"] = spec_low
    df.loc[both_na, "규격상한"] = spec_high
    col_map = {"batch": "제조번호", "test": "시험항목", "major": "대분류",
               "time": "시점(개월)", "assay": "결과값",
               "spec_low": "규격하한", "spec_high": "규격상한"}
    rp = analyze_dataframe(df, col_map, spec_low, spec_high, "pooled")
    ri = analyze_dataframe(df, col_map, spec_low, spec_high, "independent")

    def _srow(r):
        return {"component": r.test_item, "batch": r.batch, "n": int(r.n),
                "slope": None if np.isnan(r.slope) else round(float(r.slope), 4),
                "intercept": None if np.isnan(r.intercept) else round(float(r.intercept), 3),
                "r2": None if np.isnan(r.r_squared) else round(float(r.r_squared), 4),
                "trend": r.trend,
                "shelf": None if r.shelf_life is None else round(float(r.shelf_life), 1),
                "limit": r.limiting_side}
    summary_pooled = [_srow(r) for r in sorted(rp, key=lambda r: (r.test_item, r.batch))]
    summary_indep = [_srow(r) for r in sorted(ri, key=lambda r: (r.test_item, r.batch))]
    pmap = {(r.test_item, r.batch): r for r in rp}
    imap = {(r.test_item, r.batch): r for r in ri}
    comparison = []
    for k in sorted(pmap, key=lambda k: (k[0], k[1])):
        sp = pmap[k].shelf_life
        si = imap[k].shelf_life if k in imap else None
        comparison.append({"component": k[0], "batch": k[1],
                           "pooled": None if sp is None else round(float(sp), 1),
                           "indep": None if si is None else round(float(si), 1),
                           "diff": None if (sp is None or si is None) else round(float(sp - si), 1)})
    raw = []
    for r in df.sort_values(["시험항목", "제조번호", "시점(개월)"]).to_dict("records"):
        raw.append({"제조번호": r["제조번호"], "시험항목": r["시험항목"], "대분류": r["대분류"],
                    "시점": r["시점(개월)"], "결과값": r["결과값"],
                    "규격하한": None if (isinstance(r["규격하한"], float) and np.isnan(r["규격하한"])) else r["규격하한"],
                    "규격상한": None if (isinstance(r["규격상한"], float) and np.isnan(r["규격상한"])) else r["규격상한"]})
    wp = min([r.shelf_life for r in rp if r.shelf_life is not None], default=None)
    wi = min([r.shelf_life for r in ri if r.shelf_life is not None], default=None)
    return {"ok": True, "summaryPooled": summary_pooled, "summaryIndep": summary_indep,
            "comparison": comparison, "raw": raw,
            "worstPooled": None if wp is None else round(float(wp), 1),
            "worstIndep": None if wi is None else round(float(wi), 1)}


def stability_excel(code: str, test_type: str, spec_low: float = 90.0,
                    spec_high: float = 150.0, batch: Optional[str] = None) -> bytes:
    """안정성 분석결과 Excel(.xlsx) — 두방식비교·요약(통합/독립)·원자료 4개 시트."""
    from openpyxl import Workbook
    t = stability_tables(code, test_type, spec_low, spec_high, batch)
    wb = Workbook()
    ws = wb.active
    ws.title = "두방식비교"
    ws.append(["시험항목", "제조번호", "통합(권장)", "배치별독립", "차이(통합-독립)"])
    for r in t.get("comparison", []):
        ws.append([r["component"], r["batch"], r["pooled"], r["indep"], r["diff"]])
    for title, key in [("요약_통합", "summaryPooled"), ("요약_독립", "summaryIndep")]:
        s = wb.create_sheet(title)
        s.append(["시험항목", "제조번호", "데이터수", "기울기", "절편", "R2", "추세", "유효기간(개월)", "제한규격"])
        for r in t.get(key, []):
            s.append([r["component"], r["batch"], r["n"], r["slope"], r["intercept"],
                      r["r2"], r["trend"], r["shelf"], r["limit"]])
    s = wb.create_sheet("원자료")
    s.append(["제조번호", "시험항목", "대분류", "시점(개월)", "결과값", "규격하한", "규격상한"])
    for r in t.get("raw", []):
        s.append([r["제조번호"], r["시험항목"], r["대분류"], r["시점"], r["결과값"], r["규격하한"], r["규격상한"]])
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
